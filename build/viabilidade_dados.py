# Lê HelpCure_Simulador_Vidas.xlsx (já recalculada) e gera assets/viabilidade-dados.js para a página do investidor.
# Uso: python3 build/viabilidade_dados.py <HelpCure_Simulador_Vidas.xlsx>
import json, sys, os, datetime
from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
wb = load_workbook(sys.argv[1], data_only=True)
b, p = wb["Base_Escopos"], wb["Premissas"]
TIERS = ["PY", "BR base", "BR alto", "BR realista"]
CODES = ["E1", "E2", "E3", "E4", "E5"]
REV0, SRV0, COP0, UOP0, CPX0, LIV0 = 70, 93, 101, 109, 117, 125   # layout fixo do build_simulador_vidas.py
row = lambda r: [b.cell(r, c).value or 0 for c in range(2, 8)]
P = lambda a: p[a].value

# Pelo menos 8 meses de desenvolvimento/implantação (Ano 0, sem receita) em todos os segmentos — decisão de 30/10/2026.
MIN_DEV_MESES = 8
esc = {}
for i, code in enumerate(CODES):
    r = 5 + i
    v = [b.cell(r, c).value for c in range(1, 19)]
    tier = v[8]
    esc[code] = {
        "nbase": v[2], "m0": max(v[3], MIN_DEV_MESES), "teamb": v[4], "integ": v[5], "roy": v[6], "regime": v[7], "tabela": tier,
        "passo": v[11], "min": v[12], "max": v[13], "celula": v[14], "rodada": v[15], "elast": v[16], "integk": v[17],
        "vidas": row(LIV0 + i), "receita": row(REV0 + i * 4 + TIERS.index(tier)),
        "servir": row(SRV0 + i), "op_celula": row(COP0 + i), "op_movel": row(UOP0 + i), "capex": row(CPX0 + i),
    }
prem = {
    "mkt": P("B15"), "grow": P("B16"), "safety": P("B17"), "taxPY": P("B18"), "taxBR": P("B19"), "wPY": P("B20"), "wBR": P("B21"),
    "floor": P("B23"), "infra": P("B24"), "dev": P("B27"), "pil": P("B28"), "cg": P("B29"), "corp": P("B30"), "f1cap": P("B31"),
    "tr": [P("B32"), P("B33"), P("B34")], "tir": P("B37"), "mult": P("B38"), "corpk": P("B39"), "res": P("B40"), "payout": P("B41"),
    "minm": P("B42"), "eref": P("B43"), "dsh": P("B44"), "psh": P("B45"), "csh": P("B46"),
}
# ---- Clínica (30/10/2026): como fundadora, paga a Fase 1 cheia e recebe a cota fundadora. A Fase 1 enxuta abaixo ficou só como referência.
# Antes: projeto enxuto, em duas versões — Paraguai e Brasil (preços e impostos de cada país).
# Só desenvolvimento + implantação + 24 meses de manutenção da empresa. Sem dispositivos, célula, unidade móvel nem piloto físico.
# Itens da aba F1_Desenvolvimento (HelpCure_Modelo_por_Abrangencia4.xlsx, pontos médios) que saem ou diminuem:
F1_DEV_SUB = 2248000                 # subtotal do desenvolvimento completo
F1_FORA = {
    "Software das células e unidades móveis": 120000,
    "Engenharia da unidade móvel": 60000,
    "Ingestão de dispositivos: metade (sem kit Help Cure Casa nem SDKs OEM; fica o Health Connect)": 100000,
}
CONT = 0.10
dev_cli = (F1_DEV_SUB - sum(F1_FORA.values())) * (1 + CONT)
cg = prem["cg"]                       # manutenção da empresa por 24 meses (igual à Fase 1 completa)
f1_cli = {"dev": dev_cli, "pil": 0, "cg": cg, "f1cap": (dev_cli + cg) * prem["safety"]}
base = esc.pop("E1")
prem_elast = 0.7   # suporte com ganho de escala, como nos demais segmentos (decisão de 30/10/2026; a planilha usa 1,0)
for code, tier, regime in (("E1PY", "PY", "PY"), ("E1BR", "BR realista", "BR")):
    e = dict(base); e.update({"tabela": tier, "regime": regime, "rodada": 1, "elast": prem_elast,  # Fase 1 cheia (cota fundadora, 30/10/2026); f1_cli fica só como referência
                              "receita": row(REV0 + 0 * 4 + TIERS.index(tier))})
    esc[code] = e
# ---- Clínica (30/10/2026, pesquisa de mercado): 2.000 vidas monitoradas por clínica no Ano 5, crescendo 10% ao ano.
# Base: ~67 mil pacientes ativos por clínica × 3% de adesão (AmorSaúde ~19 mil/ano, dr.consulta ~100 mil, rede do Rio ~182 mil).
# Crescimento: Rede D'Or +10% em volume (3T25), Hapvida +6% em vidas, planos +2,8%, monitoramento remoto +12%/ano.
CLI_VIDAS5, CLI_CRESC, CLI_FEE = 2000, 0.10, 5500
MIX = {"Usuário": 0.10, "Governo": 0.30, "Plano": 0.60}           # mesmo mix do plano original (50/150/300 em 500)
PRECO = {"PY": {"Usuário": 9.9, "Governo": 10, "Plano": 12}, "BR realista": {"Usuário": 9.9, "Governo": 8, "Plano": 14}}
SERVIR = {"Usuário": 0.6, "Governo": 4.6, "Plano": 4.6}            # custo clínico + nuvem por vida/mês (aba Premissas)
vid = [0] + [CLI_VIDAS5 / (1 + CLI_CRESC) ** (5 - y) for y in range(1, 6)]
for code in ("E1PY", "E1BR"):
    t = esc[code]["tabela"]
    esc[code].update({
        "nbase": CLI_VIDAS5, "passo": CLI_VIDAS5, "min": CLI_VIDAS5, "max": 100000,
        "vidas": [round(v) for v in vid],
        "receita": [0] + [CLI_FEE * 12 + sum(v * MIX[k] * PRECO[t][k] * 12 for k in MIX) for v in vid[1:]],
        "servir": [0] + [sum(v * MIX[k] * SERVIR[k] * 12 for k in MIX) for v in vid[1:]],
        "op_celula": [0] * 6, "op_movel": [0] * 6, "capex": [0] * 6})
# ---- Vidas por estimativa de mercado (30/10/2026): 2% da população sem plano (SUS) + 2,4% dos beneficiários de planos no Ano 5,
# mais a parcela de empregadores do plano original. Rampa: a forma da curva original de cada pagador.
# Fontes: IBGE, estimativa de 1/7/2025; ANS/IESS NAB 111 (set/2025); Paraguai: censo 2022 (INE) e ~7% com seguro privado (Cimap).
MERCADO = {"E3": {"uf": {"SP": (46081801, 18694732), "MG": (21393441, 5902946), "RJ": (17223547, 5607424),
                         "BA": (14870907, 1791085), "PR": (11890517, 3202576), "CE": (9268836, 1489455)}},
           "E5": {"uf": {"BR": (213421037, 53232560)}},
           "E4": {"uf": {"PY": (6109644, round(6109644 * 0.07))}}}
SUS_ADES, PLANO_ADES = 0.02, 0.024
PASSO = {"E3": 20000, "E4": 20000, "E5": 100000}
pm = wb["Premissas"]
PRICE = {t: [pm.cell(5 + i, 2 + j).value for i in range(4)] for j, t in enumerate(TIERS)}   # Usuário, Governo, Plano, Empregador
COST = [(pm.cell(5 + i, 6).value or 0) + (pm.cell(5 + i, 7).value or 0) for i in range(4)]
CELCAP, CELOP, UMCAP, UMOP = (pm[c].value for c in ("B11", "B12", "B13", "B14"))
for code, cfg in MERCADO.items():
    i = CODES.index(code); bk = 12 + i * 11; r = 5 + i
    pay = [[b.cell(bk + 2 + q, 2 + y).value or 0 for y in range(6)] for q in range(4)]          # vidas por pagador, Ano 0..5
    ent = [b.cell(bk + 7, 2 + y).value or 0 for y in range(6)]; um = [b.cell(bk + 8, 2 + y).value or 0 for y in range(6)]
    fee, setup, cell = b.cell(r, 10).value, b.cell(r, 11).value, b.cell(r, 15).value
    S = esc[code]; t = S["tabela"]
    def arrays(pay, ent, um):
        rec, srv, cop, uop, cpx = [], [], [], [], []
        for y in range(6):
            pe = ent[y - 1] if y else 0; pu = um[y - 1] if y else 0
            rec.append(sum(pay[q][y] * PRICE[t][q] * 12 for q in range(4)) + ent[y] * fee * 12 + max(0, ent[y] - pe) * CELCAP * setup * cell)
            srv.append(sum(pay[q][y] * COST[q] for q in range(4)) * 12)
            cop.append(ent[y] * CELOP * 12 * cell); uop.append(um[y] * UMOP * 12)
            cpx.append(max(0, ent[y] - pe) * CELCAP * cell + max(0, um[y] - pu) * UMCAP)
        return rec, srv, cop, uop, cpx
    chk = arrays(pay, ent, um)                       # conferência: refaz as séries do plano original
    assert all(abs(x - y) < 1e-6 for a_, b_ in zip(chk, (S["receita"], S["servir"], S["op_celula"], S["op_movel"], S["capex"])) for x, y in zip(a_, b_)), code
    y5 = [pay[q][5] for q in range(4)]; tot5 = sum(y5); gov_share = y5[1] / (y5[0] + y5[1]); emp_share = y5[3] / tot5
    novo = [0, 0, 0, 0]
    for pop, plano in cfg["uf"].values():
        sus = (pop - plano) * SUS_ADES; pl = plano * PLANO_ADES; tot = (sus + pl) / (1 - emp_share)
        for q, v in enumerate((sus * (1 - gov_share), sus * gov_share, pl, tot * emp_share)): novo[q] += v / len(cfg["uf"])
    pay2 = [[pay[q][y] / y5[q] * novo[q] if y5[q] else 0 for y in range(6)] for q in range(4)]
    kt = sum(novo) / tot5
    rec, srv, cop, uop, cpx = arrays(pay2, [x * kt for x in ent], [x * kt for x in um])
    nb = round(sum(novo))
    S.update({"nbase": nb, "teamb": S["teamb"] * kt ** S["elast"], "vidas": [sum(pay2[q][y] for q in range(4)) for y in range(6)],
              "receita": rec, "servir": srv, "op_celula": cop, "op_movel": uop, "capex": cpx,
              "passo": PASSO[code], "min": PASSO[code], "max": max(S["max"], 3 * nb), "mercado": {"vidas_ano5": nb, "fator_vs_plano": kt}})

# ---- Plano de saúde (E2) por estimativa de mercado, com crianças (30/10/2026): ver build/mercado_e2.py
import sys as _s; _s.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from mercado_e2 import aplica_e2
aplica_e2(esc["E2"])
from dispositivos_modelo import aplica_disp   # dispositivos no pacote (receita = custo, preço de lote)
aplica_disp(esc)

# ---- Modelo final (30/10/2026): royalty RAPP por vida e divisão de lucros do fundador (planilha HelpCure_Equity_Lucros_Royalties)
prem.update({"mvida": 0.60, "ps": 0.10, "pscap": 5, "psres": 0.01})
for _c in ("E1PY", "E1BR"):
    esc[_c]["ref"] = 60000  # referência da barra: rede de 30 clínicas
    esc[_c]["psmax"] = 0.35  # rede de clínicas fundadora: divisão de lucros de 10% a 35%, para vender no máximo 30%
out = {"fonte": "HelpCure_Simulador_Vidas.xlsx", "gerado": datetime.date.today().isoformat(), "premissas": prem, "escopos": esc}
dst = os.path.join(ROOT, "assets", "viabilidade-dados.js")
open(dst, "w", encoding="utf-8").write("// Gerado por build/viabilidade_dados.py a partir do simulador de vidas. Não editar à mão.\nwindow.HC_VIAB=" + json.dumps(out, ensure_ascii=False, separators=(",", ":")) + ";\n")
print("ok", dst, os.path.getsize(dst), "bytes")
